import time
import os
import pandas as pd
import numpy as np
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
import openpyxl
from openpyxl.styles import PatternFill, Font

INPUT_DIR = r"D:\14. Project\04. Data Cleaning\DIRTY"
OUTPUT_DIR = r"D:\14. Project\04. Data Cleaning\CLEANED"

def clean_mining_data(file_path):
    print(f"\n[INFO] Memproses file: {os.path.basename(file_path)}")
    start_time = time.time()
    
    try:
        print("[STEP 1/7] Membaca file Excel...")
        if file_path.endswith('.xlsx') or file_path.endswith('.xls'):
            df = pd.read_excel(file_path)
        elif file_path.endswith('.csv'):
            df = pd.read_csv(file_path)
        else:
            print("[WARNING] Format file tidak didukung.")
            return
    except Exception as e:
        print(f"[ERROR] Gagal membaca file: {e}")
        return

    initial_rows = len(df)
    print(f"[INFO] Total baris awal: {initial_rows:,} baris | Kolom: {len(df.columns)}")

    print("[STEP 2/7] Standardisasi nama kolom & ubah 'unknown' menjadi NaN...")
    df.columns = df.columns.str.strip().str.lower().str.replace(' ', '_')
    df.replace(['-', 'TBD', 'tbd', 'unknown', 'UNKNOWN', 'Unknown', '?', 'N/A', 'nan', ''], np.nan, inplace=True)
    
    text_columns = df.select_dtypes(include=['object', 'string']).columns
    for col in text_columns:
        df[col] = df[col].astype(str).str.strip()
        df[col] = df[col].replace(['nan', 'UNKNOWN', 'unknown', 'Unknown'], np.nan)
        
    df = df.drop_duplicates()
    df = df.reset_index(drop=True)

    print("[STEP 3/7] Menyamakan tahun pada kolom 'date' menjadi 2025...")
    if 'date' in df.columns:
        df['date'] = pd.to_datetime(df['date'], errors='coerce')
        df['date'] = df['date'].apply(lambda x: x.replace(year=2025) if pd.notnull(x) else x)

    print("[STEP 4/7] Filter Baris: Hapus baris jika >50% sel kosong...")
    threshold_empty = len(df.columns) * 0.5
    before_row_filter = len(df)
    df = df.dropna(thresh=len(df.columns) - threshold_empty)
    df = df.reset_index(drop=True)
    print(f"[INFO] Menghapus {before_row_filter - len(df):,} baris karena terlalu banyak sel kosong (>50%).")

    print("[STEP 5/7] Validasi Rentang Fisik Logis (Ubah data di luar batas menjadi NaN)...")
    valid_ranges = {
        'est': (0, 30),
        'fuel_burn': (0, 150),
        'eco_mode': (0, 100),
        'payload': (0, 130)
    }
    
    for col, (min_val, max_val) in valid_ranges.items():
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='coerce')
            invalid_mask = (df[col] < min_val) | (df[col] > max_val)
            count_invalid = invalid_mask.sum()
            df.loc[invalid_mask, col] = np.nan
            print(f" - Kolom {col}: Mengubah {count_invalid:,} data di luar rentang ({min_val} - {max_val}) menjadi NaN.")

    print("[STEP 6/7] Deteksi Outlier metode IQR pada data valid & Imputasi Missing Value...")
    outlier_cols = ['fuel_burn', 'eco_mode', 'payload', 'est']
    outlier_indices = {}
    
    for col in outlier_cols:
        if col in df.columns:
            series_clean = df[col].dropna()
            if not series_clean.empty:
                Q1 = series_clean.quantile(0.25)
                Q3 = series_clean.quantile(0.75)
                IQR = Q3 - Q1
                lower_bound = Q1 - 1.5 * IQR
                upper_bound = Q3 + 1.5 * IQR
                
                is_outlier = (df[col].notnull()) & ((df[col] < lower_bound) | (df[col] > upper_bound))
                outlier_indices[col] = df[is_outlier].index.tolist()
                print(f" - IQR {col}: Batas Bawah={lower_bound:.2f}, Batas Atas={upper_bound:.2f} ({len(outlier_indices[col]):,} outlier terdeteksi)")

    df = df.ffill().bfill()

    for col in ['est', 'fuel_burn', 'eco_mode', 'payload']:
        if col in df.columns:
            if df[col].isnull().sum() > 0:
                med_val = df[col].median()
                df[col] = df[col].fillna(med_val)

    cat_cols = df.select_dtypes(include=['object']).columns
    for col in cat_cols:
        if df[col].isnull().sum() > 0:
            mode_val = df[col].mode()[0] if not df[col].mode().empty else "Unknown"
            df[col] = df[col].fillna(mode_val)

    if 'blast' in df.columns:
        df['blast'] = df['blast'].astype(str).str.upper()
    if 'dump' in df.columns:
        df['dump'] = df['dump'].astype(str).str.upper()
    if 'jalan' in df.columns:
        df['jalan'] = df['jalan'].astype(str).str.title()

    print("[STEP 7/7] Menyimpan file Excel & Mewarnai Outlier menjadi merah...")
    file_name = os.path.basename(file_path)
    output_path = os.path.join(OUTPUT_DIR, f"cleaned_{file_name}")
    
    df.to_excel(output_path, index=False)
    
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="9C0006", bold=True)
    
    col_name_to_idx = {cell.value: idx for idx, cell in enumerate(ws[1], 1)}
    
    for col_name, indices in outlier_indices.items():
        if col_name in col_name_to_idx:
            c_idx = col_name_to_idx[col_name]
            for row_idx in indices:
                excel_row = row_idx + 2
                cell = ws.cell(row=excel_row, column=c_idx)
                cell.fill = red_fill
                cell.font = red_font
                
    wb.save(output_path)
    
    elapsed_time = time.time() - start_time
    print(f"[SUCCESS] Selesai dalam {elapsed_time:.2f} detik!")
    print(f"Baris Awal: {initial_rows:,} | Baris Akhir Bersih: {len(df):,}")
    print(f"File bersih tersimpan di: {output_path}\n")

class FolderHandler(FileSystemEventHandler):
    def on_created(self, event):
        if event.is_directory:
            return
        if event.src_path.endswith(('.xlsx', '.xls', '.csv')):
            file_size = 0
            while file_size == 0:
                try:
                    file_size = os.path.getsize(event.src_path)
                except OSError:
                    time.sleep(1)
                if file_size == 0:
                    time.sleep(1)
            time.sleep(2)
            clean_mining_data(event.src_path)

if __name__ == "__main__":
    os.makedirs(INPUT_DIR, exist_ok=True)
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    
    print("==================================================")
    print("  AUTOMATED FMS FOLDER WATCHER & CLEANER AKTIF  ")
    print("==================================================")
    print(f"Mengawasi folder DIRTY : {INPUT_DIR}")
    print(f"Output folder CLEANED  : {OUTPUT_DIR}")
    print("Tekan Ctrl + C di terminal untuk berhenti.\n")
    
    event_handler = FolderHandler()
    observer = Observer()
    observer.schedule(event_handler, path=INPUT_DIR, recursive=False)
    observer.start()
    
    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()
        print("\n[INFO] Program dihentikan.")
    observer.join()